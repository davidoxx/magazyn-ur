"""Zamienia arkusz części (Excel) w katalog wbudowany w aplikację.

Użycie (z katalogu projektu):
    python narzedzia/katalog_z_excela.py "C:/sciezka/Czesci zam 2508.xlsx"

Wynik: shared/src/commonMain/composeResources/files/katalog.json (wspólny dla aplikacji
na Androida i programu na Windows). Potem zbuduj obie.

Arkusze:
- "Data": kolumna A pełny numer SAP, B krótki tekst z SAP, C miejsce składowania.
  Tylko te pozycje trafiają do katalogu.
- "Magazyn*" (Magazyn1, Magazyn2, ...): A końcówka numeru SAP, B nazwa, C dostawca.
  Pełny numer to "538" + końcówka dopełniona zerami do 12 znaków (1000 -> 538000001000).
  Nazwy z tych arkuszy są pełniejsze, więc służą do wyszukiwania i są pokazywane
  obok nazwy z SAP. Końcówek, których nie ma w arkuszu Data, nie dodajemy.
"""

import datetime
import json
import re
import sys
from pathlib import Path

import openpyxl

WYNIK = (Path(__file__).resolve().parent.parent / "shared" / "src" / "commonMain"
         / "composeResources" / "files" / "katalog.json")
PELNY_SAP = re.compile(r"538\d{9}")


def tekst(wartosc) -> str:
    """Komórka jako tekst: bez nowych linii i wielokrotnych spacji."""
    if wartosc is None:
        return ""
    return re.sub(r"\s+", " ", str(wartosc)).strip()


def numer_sap(wartosc) -> str:
    s = tekst(wartosc)
    # Excel zapisał część numerów jako liczby: "5380011714.0"
    if re.fullmatch(r"\d+\.0", s):
        s = s[:-2]
    return s


def miejsca(wartosc) -> list[str]:
    # Część miejsc ma śmieci z eksportu: ";22A;;;;;;", ";S13D"
    return [m.strip() for m in tekst(wartosc).split(";") if m.strip()]


def dodaj_unikalne(lista: list[str], wartosc: str) -> None:
    if wartosc and wartosc.casefold() not in (x.casefold() for x in lista):
        lista.append(wartosc)


def main() -> None:
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    plik = Path(sys.argv[1])
    wb = openpyxl.load_workbook(plik, read_only=True, data_only=True)
    if "Data" not in wb.sheetnames:
        sys.exit('Brak arkusza "Data"')

    pozycje: dict[str, dict] = {}
    for wiersz in list(wb["Data"].iter_rows(values_only=True))[1:]:
        sap = numer_sap(wiersz[0] if len(wiersz) > 0 else None)
        if not sap:
            continue
        p = pozycje.setdefault(sap, {"sap": sap, "nazwa": "", "miejsca": [], "inneNazwy": [], "dostawcy": []})
        # Ten sam numer bywa w kilku wierszach: zwykle drugi bez miejsca, czasem inne miejsce
        if not p["nazwa"]:
            p["nazwa"] = tekst(wiersz[1] if len(wiersz) > 1 else None)
        for m in miejsca(wiersz[2] if len(wiersz) > 2 else None):
            dodaj_unikalne(p["miejsca"], m)

    pominiete = 0
    arkusze_magazyn = [n for n in wb.sheetnames if n.lower().startswith("magazyn")]
    for nazwa_arkusza in arkusze_magazyn:
        for wiersz in list(wb[nazwa_arkusza].iter_rows(values_only=True))[1:]:
            koncowka = numer_sap(wiersz[0] if len(wiersz) > 0 else None)
            if not koncowka.isdigit():
                continue
            p = pozycje.get("538" + koncowka.zfill(9))
            if p is None:
                pominiete += 1
                continue
            nazwa = tekst(wiersz[1] if len(wiersz) > 1 else None)
            if nazwa.casefold() != p["nazwa"].casefold():
                dodaj_unikalne(p["inneNazwy"], nazwa)
            dodaj_unikalne(p["dostawcy"], tekst(wiersz[2] if len(wiersz) > 2 else None))

    # Puste listy pomijamy, aplikacja ma dla nich wartości domyślne
    lista = [{k: v for k, v in p.items() if v} for p in pozycje.values()]
    katalog = {
        "zrodlo": plik.name,
        "utworzono": datetime.date.today().isoformat(),
        "pozycje": lista,
    }
    WYNIK.parent.mkdir(parents=True, exist_ok=True)
    WYNIK.write_text(json.dumps(katalog, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")

    print(f"Zapisano {WYNIK}")
    print(f"  pozycji: {len(lista)}")
    print(f"  z innym formatem numeru niż 538 + 9 cyfr: {sum(1 for p in lista if not PELNY_SAP.fullmatch(p['sap']))}")
    print(f"  bez miejsca: {sum(1 for p in lista if not p.get('miejsca'))}")
    print(f"  z kilkoma miejscami: {sum(1 for p in lista if len(p.get('miejsca', [])) > 1)}")
    print(f"  z nazwą z arkuszy {', '.join(arkusze_magazyn)}: {sum(1 for p in lista if p.get('inneNazwy'))}")
    print(f"  pominięte wiersze z arkuszy Magazyn (brak numeru w Data): {pominiete}")


if __name__ == "__main__":
    main()
